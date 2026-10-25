"""Bounded local cost-plan extraction. All output remains an editable estimate draft."""
import csv
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import re
import zipfile

MAX_ROWS = 500
MAX_TEXT = 100000
FIELDS = ('title', 'category', 'estimate', 'kind', 'company_hint', 'notes', 'is_group')


def is_summary(title):
    return bool(re.search(r'(?i)\b(gesamt(?:summe|kosten|betrag)?|summe|zwischensumme|übertrag)\b', title)
                or re.match(r'(?i)^(netto(?:betrag|summe)?|brutto(?:betrag|summe)?|umsatzsteuer|mehrwertsteuer|mwst\.?|ust\.?)\b', title))


def amount(value):
    if isinstance(value, bool) or value is None:
        raise ValueError('Schätzung fehlt.')
    if isinstance(value, str):
        value = value.strip().replace('€', '').replace('EUR', '').replace('\u00a0', '').replace(' ', '')
        if ',' in value:
            value = value.replace('.', '').replace(',', '.')
        elif re.fullmatch(r'\d{1,3}(\.\d{3})+', value):
            value = value.replace('.', '')
    try:
        result = Decimal(str(value))
        if not result.is_finite() or result < 0 or result > 999999999 or result.as_tuple().exponent < -2:
            raise ValueError('Ungültige Schätzung (maximal zwei Nachkommastellen).')
        return float(result)
    except (InvalidOperation, TypeError):
        raise ValueError('Schätzung ist kein gültiger Eurobetrag.')


def checked_rows(rows):
    if not isinstance(rows, list) or not 1 <= len(rows) <= MAX_ROWS:
        raise ValueError('Bitte eine bis 500 Planpositionen prüfen.')
    result = []
    for number, raw in enumerate(rows, 1):
        if not isinstance(raw, dict):
            raise ValueError('Ungültige Planposition.')
        row = {}
        for key, limit in [('title', 300), ('category', 160), ('company_hint', 200), ('notes', 2000)]:
            value = raw.get(key, '')
            if not isinstance(value, str) or len(value) > limit:
                raise ValueError(f'Zeile {number}: {key} ist zu lang oder ungültig.')
            row[key] = value.strip()
        if not row['title'] or not row['category']:
            raise ValueError(f'Zeile {number}: Position und Gewerk fehlen.')
        row['kind'] = raw.get('kind', 'Fremdleistung')
        if row['kind'] not in ('Fremdleistung', 'Eigenleistung'):
            raise ValueError('Ungültige Leistungsart.')
        if type(raw.get('is_group', False)) is not bool:
            raise ValueError('Gruppenkennzeichnung muss wahr oder falsch sein.')
        row['is_group'] = raw.get('is_group', False)
        row['estimate'] = amount(raw.get('estimate'))
        result.append(row)
    return result


ALIASES = {
    'title': ('position', 'bezeichnung', 'beschreibung', 'leistung', 'title', 'gewerk / position'),
    'category': ('gewerk', 'kategorie', 'kostengruppe', 'category'),
    'estimate': ('schätzung', 'schatzung', 'schätzung (€)', 'gesamtpreis', 'gesamtbetrag', 'kosten', 'betrag', 'brutto', 'estimate'),
    'kind': ('art', 'leistungsart', 'kind'),
    'company_hint': ('firma', 'firmenhinweis', 'anbieter', 'company_hint'),
    'notes': ('notiz', 'notizen', 'bemerkung', 'notes'),
}


def _table_rows(table):
    mapping = {}
    result = []
    for cells in table:
        normalized = [str(c or '').strip().lower() for c in cells]
        candidate = {field: i for i, cell in enumerate(normalized)
                     for field, names in ALIASES.items() if cell in names}
        if 'title' in candidate and 'estimate' in candidate:
            mapping = candidate
            continue
        if not mapping:
            continue
        row = {field: cells[i] if i < len(cells) else '' for field, i in mapping.items()}
        if not str(row.get('title', '') or '').strip():
            continue
        try:
            row['estimate'] = amount(row.get('estimate'))
        except ValueError:
            continue
        for field in ('title', 'category', 'company_hint', 'notes'):
            row[field] = str(row.get(field, '') or '').strip()
        row['category'] = row['category'] or 'Sonstiges'
        row['kind'] = 'Eigenleistung' if str(row.get('kind', '')).lower() == 'eigenleistung' else 'Fremdleistung'
        row['is_group'] = is_summary(row['title'])
        result.append(row)
        if len(result) > MAX_ROWS:
            raise ValueError('Mehr als 500 Positionen. Bitte die Datei aufteilen.')
    return result


def local_plan(data, filename):
    extension = filename.rsplit('.', 1)[-1].lower()
    if extension == 'xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError('Entpackte Excel-Datei ist zu groß.')
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            rows = []
            text_parts = []
            for sheet in book.worksheets[:10]:
                if sheet.max_row > 5000 or sheet.max_column > 100:
                    raise ValueError('Excel-Tabelle ist zu groß (max. 5000 Zeilen und 100 Spalten).')
                table = list(sheet.iter_rows(values_only=True))
                rows.extend(_table_rows(table))
                if sum(map(len, text_parts)) < MAX_TEXT:
                    text_parts.append(sheet.title + '\n' + '\n'.join('\t'.join(str(c or '') for c in cells) for cells in table)[:MAX_TEXT])
        finally:
            book.close()
        text = '\n'.join(text_parts)[:MAX_TEXT]
    elif extension == 'csv':
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('cp1252')
        try:
            dialect = csv.Sniffer().sniff(text[:8000], delimiters=';,\t')
        except csv.Error:
            dialect = csv.excel
        rows = _table_rows(csv.reader(StringIO(text), dialect))
    else:
        if extension == 'txt':
            text = data.decode('utf-8-sig')
        elif extension == 'docx':
            from project_documents import text_from_document
            text, _ = text_from_document(data, '.docx')
        else:
            from receipt_parser import extract_text
            text, _ = extract_text(data, filename)
        rows = []
        for line in text[:MAX_TEXT].splitlines():
            match = re.match(r'^(.+?)\s+((?:\d{1,3}(?:[. ]\d{3})+|\d+),\d{2})\s*(?:€|EUR)?\s*$', line.strip())
            if match:
                title = match[1].strip(' .;\t')[:300]
                rows.append(dict(title=title, category='Sonstiges', estimate=amount(match[2]),
                                 is_group=is_summary(title)))
    if not rows:
        return {'rows': [], 'text': text[:MAX_TEXT], 'warnings': [
            'Keine eindeutigen Positionen erkannt. Positionen manuell ergänzen oder ChatGPT-Analyse wählen.']}
    return {'rows': checked_rows(rows), 'text': text[:MAX_TEXT], 'warnings': [
        'Vollständigkeit, Brutto/Netto und Summenzeilen mit dem Original prüfen. Gruppen werden nicht summiert. '
        'Excel-Formeln benötigen gespeicherte Ergebniswerte; nicht lesbare Zeilen können fehlen.']}
